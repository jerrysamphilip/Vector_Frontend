# app/services/reports_service.py
"""
Reports Service — aggregation queries for Global Analytics, Campaign Comparison,
Provider Performance, and CSV exports.
"""

import io
import csv
import logging
from datetime import date, timedelta, datetime
from typing import Optional

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from sqlalchemy.orm import Session
from sqlalchemy import func, and_, case

from app.models.campaign import Campaign, CampaignProspect
from app.models.metrics import CampaignMetricsRealtime
from app.models.email_message import EmailMessage, EmailEvent
from app.models.sending_inbox import SendingInbox
from app.models.prospect import Prospect
from app.models.domain_reputation import SendingDomain, DomainHealthSnapshot
from app.models.user import User
from app.schemas.reports_schema import (
    GlobalKPIs,
    TimeSeriesPoint,
    GlobalAnalyticsResponse,
    CampaignComparisonItem,
    CampaignComparisonResponse,
    ProviderPerformanceItem,
    ProviderPerformanceResponse,
    LeadStats,
    MailboxStats,
    CampaignStats,
    MailboxHealthItem,
    CampaignOptimizationMetrics,
    UserPerformanceItem,
    DashboardSummaryResponse,
)

logger = logging.getLogger(__name__)


class ReportsService:
    def __init__(self, db: Session):
        self.db = db

    def _campaign_subquery(self, tenant_id: str, user_id: Optional[str] = None):
        """Return a subquery of campaign_ids scoped to tenant (and optionally a single user)."""
        q = self.db.query(Campaign.campaign_id).filter(Campaign.tenant_id == tenant_id)
        if user_id:
            q = q.filter(Campaign.created_by == user_id)
        return q.subquery()

    # ------------------------------------------------------------------
    # GLOBAL ANALYTICS
    # ------------------------------------------------------------------

    def get_global_analytics(
        self,
        tenant_id: str,
        start_date: date,
        end_date: date,
        granularity: str = "daily",
        user_id: Optional[str] = None,
    ) -> GlobalAnalyticsResponse:
        kpis = self._get_global_kpis(tenant_id, start_date, end_date, user_id)
        timeseries = self._get_global_timeseries(tenant_id, start_date, end_date, granularity, user_id)
        return GlobalAnalyticsResponse(kpis=kpis, timeseries=timeseries)

    def _get_global_kpis(self, tenant_id: str, start_date: date, end_date: date, user_id: Optional[str] = None) -> GlobalKPIs:
        """Aggregate KPIs across all campaigns for a tenant (optionally scoped to a user)."""

        tenant_campaigns = self._campaign_subquery(tenant_id, user_id)

        sent = (
            self.db.query(func.count(EmailMessage.message_id))
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.status == "SENT",
                EmailMessage.sent_at >= start_date,
                EmailMessage.sent_at < end_date + timedelta(days=1),
            )
            .scalar()
        ) or 0

        event_row = (
            self.db.query(
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_OPEN, 1), else_=0)).label("opened"),
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_REPLY, 1), else_=0)).label("replied"),
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_BOUNCE, 1), else_=0)).label("bounced"),
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_CLICK, 1), else_=0)).label("clicked"),
            )
            .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailEvent.event_time >= start_date,
                EmailEvent.event_time < end_date + timedelta(days=1),
            )
            .first()
        )

        opened = event_row.opened or 0
        replied = event_row.replied or 0
        bounced = event_row.bounced or 0
        clicked = event_row.clicked or 0

        # Fallback: if no open events, count prospects with OPENED/REPLIED status
        # (mirrors CampaignAnalyticsService logic — a reply implies the email was opened)
        opened_prospect_fallback = (
            self.db.query(func.count(func.distinct(CampaignProspect.prospect_id)))
            .filter(
                CampaignProspect.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                CampaignProspect.status.in_(["OPENED", "REPLIED"]),
            )
            .scalar()
        ) or 0
        opened = max(opened, opened_prospect_fallback)

        replied_prospect_fallback = (
            self.db.query(func.count(func.distinct(CampaignProspect.prospect_id)))
            .filter(
                CampaignProspect.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                CampaignProspect.status == "REPLIED",
            )
            .scalar()
        ) or 0
        replied = max(replied, replied_prospect_fallback)

        # Leads contacted
        leads_contacted = (
            self.db.query(func.count(func.distinct(EmailMessage.prospect_id)))
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.status == "SENT",
                EmailMessage.sent_at >= start_date,
                EmailMessage.sent_at < end_date + timedelta(days=1),
            )
            .scalar()
        ) or 0

        # Active campaigns — scoped to user if provided
        active_campaigns = (
            self.db.query(func.count(Campaign.campaign_id))
            .filter(
                Campaign.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                Campaign.status == "ACTIVE",
            )
            .scalar()
        ) or 0

        return GlobalKPIs(
            total_leads_contacted=leads_contacted,
            total_emails_sent=sent,
            total_opened=opened,
            total_replied=replied,
            total_bounced=bounced,
            total_clicked=clicked,
            overall_open_rate=round((opened / sent) * 100, 2) if sent > 0 else 0.0,
            overall_reply_rate=round((replied / sent) * 100, 2) if sent > 0 else 0.0,
            overall_bounce_rate=round((bounced / sent) * 100, 2) if sent > 0 else 0.0,
            active_campaigns=active_campaigns,
        )

    def _get_global_timeseries(
        self,
        tenant_id: str,
        start_date: date,
        end_date: date,
        granularity: str,
        user_id: Optional[str] = None,
    ) -> list:
        """Daily or weekly sent/opened/replied/bounced across all campaigns."""

        tenant_campaigns = self._campaign_subquery(tenant_id, user_id)

        # Sent per day
        sent_results = (
            self.db.query(
                func.date(EmailMessage.sent_at).label("d"),
                func.count(EmailMessage.message_id).label("cnt"),
            )
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.status == "SENT",
                EmailMessage.sent_at >= start_date,
                EmailMessage.sent_at < end_date + timedelta(days=1),
            )
            .group_by(func.date(EmailMessage.sent_at))
            .all()
        )
        sent_map = {row.d: row.cnt for row in sent_results}

        # Events per day (open, reply, bounce)
        event_results = (
            self.db.query(
                func.date(EmailEvent.event_time).label("d"),
                EmailEvent.event_type.label("t"),
                func.count(EmailEvent.event_id).label("cnt"),
            )
            .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailEvent.event_time >= start_date,
                EmailEvent.event_time < end_date + timedelta(days=1),
                EmailEvent.event_type.in_([
                    EmailEvent.EVENT_OPEN,
                    EmailEvent.EVENT_REPLY,
                    EmailEvent.EVENT_BOUNCE,
                ]),
            )
            .group_by(func.date(EmailEvent.event_time), EmailEvent.event_type)
            .all()
        )

        open_map = {}
        reply_map = {}
        bounce_map = {}
        for row in event_results:
            if row.t == EmailEvent.EVENT_OPEN:
                open_map[row.d] = row.cnt
            elif row.t == EmailEvent.EVENT_REPLY:
                reply_map[row.d] = row.cnt
            elif row.t == EmailEvent.EVENT_BOUNCE:
                bounce_map[row.d] = row.cnt

        # Build series
        points = []
        current = start_date
        while current <= end_date:
            points.append(TimeSeriesPoint(
                date=current.isoformat(),
                sent=sent_map.get(current, 0),
                opened=open_map.get(current, 0),
                replied=reply_map.get(current, 0),
                bounced=bounce_map.get(current, 0),
            ))
            current += timedelta(days=1)

        # Aggregate to weekly if requested
        if granularity == "weekly" and len(points) > 7:
            weekly = []
            for i in range(0, len(points), 7):
                chunk = points[i : i + 7]
                weekly.append(TimeSeriesPoint(
                    date=chunk[0].date,
                    sent=sum(p.sent for p in chunk),
                    opened=sum(p.opened for p in chunk),
                    replied=sum(p.replied for p in chunk),
                    bounced=sum(p.bounced for p in chunk),
                ))
            return weekly

        return points

    # ------------------------------------------------------------------
    # CAMPAIGN COMPARISON
    # ------------------------------------------------------------------

    def get_campaign_comparison(
        self,
        tenant_id: str,
        sort_by: str = "open_rate",
        sort_dir: str = "desc",
        page: int = 1,
        page_size: int = 20,
        start_date: Optional[date] = None,
        end_date: Optional[date] = None,
        user_id: Optional[str] = None,
    ) -> CampaignComparisonResponse:
        """All campaigns with metrics, sortable and paginated."""

        sent_filters = [EmailMessage.status == "SENT"]
        if start_date:
            sent_filters.append(EmailMessage.sent_at >= start_date)
        if end_date:
            sent_filters.append(EmailMessage.sent_at < end_date + timedelta(days=1))

        sent_subquery = (
            self.db.query(
                EmailMessage.campaign_id.label("campaign_id"),
                func.count(EmailMessage.message_id).label("sent_count"),
            )
            .filter(*sent_filters)
            .group_by(EmailMessage.campaign_id)
            .subquery()
        )

        event_filters = []
        if start_date:
            event_filters.append(EmailEvent.event_time >= start_date)
        if end_date:
            event_filters.append(EmailEvent.event_time < end_date + timedelta(days=1))

        event_subquery = (
            self.db.query(
                EmailMessage.campaign_id.label("campaign_id"),
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_OPEN, 1), else_=0)).label("opened_count"),
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_REPLY, 1), else_=0)).label("replied_count"),
                func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_BOUNCE, 1), else_=0)).label("bounced_count"),
            )
            .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
            .filter(*event_filters)
            .group_by(EmailMessage.campaign_id)
            .subquery()
        )

        sent_count_col = func.coalesce(sent_subquery.c.sent_count, 0)
        opened_count_col = func.coalesce(event_subquery.c.opened_count, 0)
        replied_count_col = func.coalesce(event_subquery.c.replied_count, 0)
        bounced_count_col = func.coalesce(event_subquery.c.bounced_count, 0)

        query = (
            self.db.query(
                Campaign.campaign_id,
                Campaign.campaign_name,
                Campaign.status,
                Campaign.created_at,
                sent_count_col.label("sent_count"),
                opened_count_col.label("opened_count"),
                replied_count_col.label("replied_count"),
                bounced_count_col.label("bounced_count"),
            )
            .outerjoin(sent_subquery, sent_subquery.c.campaign_id == Campaign.campaign_id)
            .outerjoin(event_subquery, event_subquery.c.campaign_id == Campaign.campaign_id)
            .filter(
                Campaign.tenant_id == tenant_id,
                *([Campaign.created_by == user_id] if user_id else []),
            )
        )

        total = query.count()

        open_rate_expr = case(
            (sent_count_col > 0, (opened_count_col * 100.0) / sent_count_col),
            else_=0.0,
        )
        reply_rate_expr = case(
            (sent_count_col > 0, (replied_count_col * 100.0) / sent_count_col),
            else_=0.0,
        )
        bounce_rate_expr = case(
            (sent_count_col > 0, (bounced_count_col * 100.0) / sent_count_col),
            else_=0.0,
        )

        # Sort
        sort_column_map = {
            "sent_count": sent_count_col,
            "opened_count": opened_count_col,
            "replied_count": replied_count_col,
            "bounced_count": bounced_count_col,
            "campaign_name": Campaign.campaign_name,
            "created_at": Campaign.created_at,
        }

        if sort_by in ("open_rate", "reply_rate", "bounce_rate"):
            rate_field_map = {
                "open_rate": open_rate_expr,
                "reply_rate": reply_rate_expr,
                "bounce_rate": bounce_rate_expr,
            }
            col = func.coalesce(rate_field_map[sort_by], 0)
        else:
            col = sort_column_map.get(sort_by, Campaign.created_at)

        if sort_dir == "asc":
            query = query.order_by(col.asc())
        else:
            query = query.order_by(col.desc())

        rows = query.offset((page - 1) * page_size).limit(page_size).all()

        # Build prospect-status fallback map for all returned campaigns
        campaign_ids = [r.campaign_id for r in rows]
        prospect_fallbacks = {}
        if campaign_ids:
            fb_rows = (
                self.db.query(
                    CampaignProspect.campaign_id,
                    CampaignProspect.status,
                    func.count(func.distinct(CampaignProspect.prospect_id)).label("cnt"),
                )
                .filter(
                    CampaignProspect.campaign_id.in_(campaign_ids),
                    CampaignProspect.status.in_(["OPENED", "REPLIED"]),
                )
                .group_by(CampaignProspect.campaign_id, CampaignProspect.status)
                .all()
            )
            for fb in fb_rows:
                if fb.campaign_id not in prospect_fallbacks:
                    prospect_fallbacks[fb.campaign_id] = {"OPENED": 0, "REPLIED": 0}
                prospect_fallbacks[fb.campaign_id][fb.status] = fb.cnt

        items = []
        for r in rows:
            sent = r.sent_count or 0
            opened = r.opened_count or 0
            replied = r.replied_count or 0
            bounced = r.bounced_count or 0
            fb = prospect_fallbacks.get(r.campaign_id, {})
            opened = max(opened, fb.get("OPENED", 0) + fb.get("REPLIED", 0))
            replied = max(replied, fb.get("REPLIED", 0))
            items.append(CampaignComparisonItem(
                campaign_id=r.campaign_id,
                campaign_name=r.campaign_name,
                status=r.status,
                sent_count=sent,
                opened_count=opened,
                replied_count=replied,
                bounced_count=bounced,
                open_rate=round((opened / sent) * 100, 2) if sent > 0 else 0.0,
                reply_rate=round((replied / sent) * 100, 2) if sent > 0 else 0.0,
                bounce_rate=round((bounced / sent) * 100, 2) if sent > 0 else 0.0,
                created_at=r.created_at,
            ))

        return CampaignComparisonResponse(
            items=items,
            total=total,
            page=page,
            page_size=page_size,
        )

    # ------------------------------------------------------------------
    # PROVIDER PERFORMANCE
    # ------------------------------------------------------------------

    def get_provider_performance(
        self,
        tenant_id: str,
        start_date: date,
        end_date: date,
        user_id: Optional[str] = None,
    ) -> ProviderPerformanceResponse:
        """Per-provider (Gmail/Outlook/SES) aggregate metrics."""

        tenant_campaigns = self._campaign_subquery(tenant_id, user_id)

        # Sent per provider
        sent_rows = (
            self.db.query(
                func.coalesce(SendingInbox.provider, "unknown").label("provider"),
                func.count(EmailMessage.message_id).label("sent"),
            )
            .join(SendingInbox, SendingInbox.inbox_id == EmailMessage.inbox_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.status == "SENT",
                EmailMessage.sent_at >= start_date,
                EmailMessage.sent_at < end_date + timedelta(days=1),
            )
            .group_by(func.coalesce(SendingInbox.provider, "unknown"))
            .all()
        )

        providers_sent = {r.provider: r.sent for r in sent_rows}

        # Events per provider
        event_rows = (
            self.db.query(
                func.coalesce(SendingInbox.provider, "unknown").label("provider"),
                EmailEvent.event_type.label("etype"),
                func.count(EmailEvent.event_id).label("cnt"),
            )
            .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
            .join(SendingInbox, SendingInbox.inbox_id == EmailMessage.inbox_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailEvent.event_time >= start_date,
                EmailEvent.event_time < end_date + timedelta(days=1),
                EmailEvent.event_type.in_([
                    EmailEvent.EVENT_OPEN,
                    EmailEvent.EVENT_REPLY,
                    EmailEvent.EVENT_BOUNCE,
                ]),
            )
            .group_by(func.coalesce(SendingInbox.provider, "unknown"), EmailEvent.event_type)
            .all()
        )

        # Build lookup
        provider_events = {}
        for r in event_rows:
            if r.provider not in provider_events:
                provider_events[r.provider] = {}
            provider_events[r.provider][r.etype] = r.cnt

        all_providers = set(providers_sent.keys()) | set(provider_events.keys())

        items = []
        for prov in sorted(all_providers):
            sent = providers_sent.get(prov, 0)
            events = provider_events.get(prov, {})
            opened = events.get(EmailEvent.EVENT_OPEN, 0)
            replied = events.get(EmailEvent.EVENT_REPLY, 0)
            bounced = events.get(EmailEvent.EVENT_BOUNCE, 0)

            items.append(ProviderPerformanceItem(
                provider=prov,
                sent_count=sent,
                delivered_count=max(sent - bounced, 0),
                opened_count=opened,
                replied_count=replied,
                bounced_count=bounced,
                open_rate=round((opened / sent) * 100, 2) if sent > 0 else 0.0,
                reply_rate=round((replied / sent) * 100, 2) if sent > 0 else 0.0,
                bounce_rate=round((bounced / sent) * 100, 2) if sent > 0 else 0.0,
            ))

        return ProviderPerformanceResponse(providers=items)

    # ------------------------------------------------------------------
    # LEAD STATS
    # ------------------------------------------------------------------

    def get_lead_stats(self, tenant_id: str, user_id: Optional[str] = None) -> LeadStats:
        """New leads (step 1) vs follow-up leads (step > 1)."""
        tenant_campaigns = self._campaign_subquery(tenant_id, user_id)

        total = (
            self.db.query(func.count(CampaignProspect.id))
            .filter(CampaignProspect.campaign_id.in_(
                self.db.query(tenant_campaigns.c.campaign_id)
            ))
            .scalar()
        ) or 0

        new_leads = (
            self.db.query(func.count(CampaignProspect.id))
            .filter(
                CampaignProspect.campaign_id.in_(
                    self.db.query(tenant_campaigns.c.campaign_id)
                ),
                CampaignProspect.current_step == 1,
            )
            .scalar()
        ) or 0

        return LeadStats(
            total_leads_contacted=total,
            new_leads_reached=new_leads,
            follow_up_leads=total - new_leads,
        )

    # ------------------------------------------------------------------
    # MAILBOX STATS
    # ------------------------------------------------------------------

    def get_mailbox_stats(self, tenant_id: str) -> MailboxStats:
        """Mailbox connection statistics."""
        inboxes = (
            self.db.query(
                SendingInbox.status,
                SendingInbox.warmup_enabled,
            )
            .filter(SendingInbox.tenant_id == tenant_id)
            .all()
        )

        total = len(inboxes)
        active = sum(1 for i in inboxes if i.status == "ACTIVE")
        paused = sum(1 for i in inboxes if i.status == "PAUSED")
        no_warmup = sum(1 for i in inboxes if not i.warmup_enabled)

        return MailboxStats(
            total_connected=total,
            mailbox_in_use=active,
            disconnected=paused,
            without_warmup=no_warmup,
        )

    # ------------------------------------------------------------------
    # CAMPAIGN STATS
    # ------------------------------------------------------------------

    def get_campaign_stats(self, tenant_id: str, user_id: Optional[str] = None) -> CampaignStats:
        """Campaign status breakdown."""
        q = self.db.query(Campaign.status, func.count(Campaign.campaign_id)).filter(Campaign.tenant_id == tenant_id)
        if user_id:
            q = q.filter(Campaign.created_by == user_id)
        rows = (
            q.group_by(Campaign.status)
            .all()
        )
        counts = {r[0]: r[1] for r in rows}
        total = sum(counts.values())

        return CampaignStats(
            total_campaigns=total,
            active=counts.get("ACTIVE", 0),
            paused=counts.get("PAUSED", 0),
            drafted=counts.get("DRAFT", 0),
            completed=counts.get("COMPLETED", 0),
        )

    # ------------------------------------------------------------------
    # EMAIL HEALTH PER MAILBOX
    # ------------------------------------------------------------------

    def get_mailbox_health(self, tenant_id: str, user_id: Optional[str] = None) -> list:
        """Per-mailbox email metrics (optionally scoped to a specific user's campaigns)."""
        tenant_campaigns = self._campaign_subquery(tenant_id, user_id)

        # Sent + leads per inbox
        sent_rows = (
            self.db.query(
                SendingInbox.email_address.label("mailbox"),
                func.count(EmailMessage.message_id).label("sent"),
                func.count(func.distinct(EmailMessage.prospect_id)).label("leads"),
            )
            .join(SendingInbox, SendingInbox.inbox_id == EmailMessage.inbox_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.status == "SENT",
            )
            .group_by(SendingInbox.email_address)
            .all()
        )

        # Events per inbox
        event_rows = (
            self.db.query(
                SendingInbox.email_address.label("mailbox"),
                EmailEvent.event_type.label("etype"),
                func.count(EmailEvent.event_id).label("cnt"),
            )
            .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
            .join(SendingInbox, SendingInbox.inbox_id == EmailMessage.inbox_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailEvent.event_type.in_([
                    EmailEvent.EVENT_OPEN,
                    EmailEvent.EVENT_REPLY,
                    EmailEvent.EVENT_BOUNCE,
                ]),
            )
            .group_by(SendingInbox.email_address, EmailEvent.event_type)
            .all()
        )

        # Build lookup
        mailbox_events = {}
        for r in event_rows:
            if r.mailbox not in mailbox_events:
                mailbox_events[r.mailbox] = {}
            mailbox_events[r.mailbox][r.etype] = r.cnt

        items = []
        for r in sent_rows:
            events = mailbox_events.get(r.mailbox, {})
            sent = r.sent or 0
            opened = events.get(EmailEvent.EVENT_OPEN, 0)
            replied = events.get(EmailEvent.EVENT_REPLY, 0)
            bounced = events.get(EmailEvent.EVENT_BOUNCE, 0)

            items.append(MailboxHealthItem(
                mailbox=r.mailbox,
                lead_contacted=r.leads or 0,
                email_sent=sent,
                opened=opened,
                opened_rate=round((opened / sent) * 100, 2) if sent > 0 else 0.0,
                replied=replied,
                replied_rate=round((replied / sent) * 100, 2) if sent > 0 else 0.0,
                bounced=bounced,
                bounce_rate=round((bounced / sent) * 100, 2) if sent > 0 else 0.0,
            ))

        return sorted(items, key=lambda x: x.email_sent, reverse=True)

    # ------------------------------------------------------------------
    # CAMPAIGN OPTIMIZATION METRICS
    # ------------------------------------------------------------------

    def get_optimization_metrics(self, tenant_id: str, user_id: Optional[str] = None) -> CampaignOptimizationMetrics:
        """Key campaign optimization metrics."""
        tenant_campaigns = self._campaign_subquery(tenant_id, user_id)

        # Avg leads contacted before first reply per campaign
        # (total prospects / campaigns that got at least 1 reply)
        campaigns_with_replies = (
            self.db.query(func.count(func.distinct(EmailMessage.campaign_id)))
            .join(EmailEvent, EmailEvent.message_id == EmailMessage.message_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailEvent.event_type == EmailEvent.EVENT_REPLY,
            )
            .scalar()
        ) or 0

        total_prospects = (
            self.db.query(func.count(CampaignProspect.id))
            .filter(CampaignProspect.campaign_id.in_(
                self.db.query(tenant_campaigns.c.campaign_id)
            ))
            .scalar()
        ) or 0

        avg_leads = round(total_prospects / campaigns_with_replies, 1) if campaigns_with_replies > 0 else 0.0

        # Follow-up reply rate (replies from prospects on step > 1)
        follow_up_prospects = (
            self.db.query(CampaignProspect.prospect_id, CampaignProspect.campaign_id)
            .filter(
                CampaignProspect.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                CampaignProspect.current_step > 1,
            )
            .subquery()
        )

        follow_up_sent = (
            self.db.query(func.count(EmailMessage.message_id))
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.prospect_id.in_(
                    self.db.query(follow_up_prospects.c.prospect_id)
                ),
                EmailMessage.status == "SENT",
            )
            .scalar()
        ) or 0

        follow_up_replies = (
            self.db.query(func.count(EmailEvent.event_id))
            .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
            .filter(
                EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                EmailMessage.prospect_id.in_(
                    self.db.query(follow_up_prospects.c.prospect_id)
                ),
                EmailEvent.event_type == EmailEvent.EVENT_REPLY,
            )
            .scalar()
        ) or 0

        follow_up_rate = round((follow_up_replies / follow_up_sent) * 100, 2) if follow_up_sent > 0 else 0.0

        # Median time to first reply (in hours)
        try:
            reply_times = (
                self.db.query(
                    func.min(EmailEvent.event_time).label("first_reply"),
                    func.min(EmailMessage.sent_at).label("first_sent"),
                )
                .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
                .filter(
                    EmailMessage.campaign_id.in_(self.db.query(tenant_campaigns.c.campaign_id)),
                    EmailEvent.event_type == EmailEvent.EVENT_REPLY,
                )
                .group_by(EmailMessage.prospect_id)
                .all()
            )

            if reply_times:
                diffs = []
                for r in reply_times:
                    if r.first_reply and r.first_sent:
                        diff_hours = (r.first_reply - r.first_sent).total_seconds() / 3600
                        if diff_hours > 0:
                            diffs.append(diff_hours)
                diffs.sort()
                median_hours = diffs[len(diffs) // 2] if diffs else 0.0
            else:
                median_hours = 0.0
        except Exception:
            median_hours = 0.0

        return CampaignOptimizationMetrics(
            avg_leads_before_first_reply=avg_leads,
            follow_up_reply_rate=follow_up_rate,
            median_time_to_first_reply_hours=round(median_hours, 1),
        )

    # ------------------------------------------------------------------
    # FULL DASHBOARD SUMMARY
    # ------------------------------------------------------------------

    def get_user_breakdown(self, tenant_id: str, start_date: date, end_date: date) -> list:
        """Per-user campaign performance summary for team analytics."""
        users = (
            self.db.query(User)
            .filter(
                User.tenant_id == tenant_id,
                User.status == "ACTIVE",
                User.role != "PLATFORM_ADMIN",
            )
            .order_by(User.first_name)
            .all()
        )

        results = []
        for u in users:
            campaigns_q = self.db.query(Campaign.campaign_id).filter(
                Campaign.tenant_id == tenant_id,
                Campaign.created_by == u.user_id,
            ).subquery()

            campaigns_count = (
                self.db.query(func.count(Campaign.campaign_id))
                .filter(Campaign.tenant_id == tenant_id, Campaign.created_by == u.user_id)
                .scalar()
            ) or 0

            sent = (
                self.db.query(func.count(EmailMessage.message_id))
                .filter(
                    EmailMessage.campaign_id.in_(self.db.query(campaigns_q.c.campaign_id)),
                    EmailMessage.status == "SENT",
                    EmailMessage.sent_at >= start_date,
                    EmailMessage.sent_at < end_date + timedelta(days=1),
                )
                .scalar()
            ) or 0

            event_row = (
                self.db.query(
                    func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_OPEN, 1), else_=0)).label("opened"),
                    func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_REPLY, 1), else_=0)).label("replied"),
                    func.sum(case((EmailEvent.event_type == EmailEvent.EVENT_BOUNCE, 1), else_=0)).label("bounced"),
                )
                .join(EmailMessage, EmailMessage.message_id == EmailEvent.message_id)
                .filter(
                    EmailMessage.campaign_id.in_(self.db.query(campaigns_q.c.campaign_id)),
                    EmailEvent.event_time >= start_date,
                    EmailEvent.event_time < end_date + timedelta(days=1),
                )
                .first()
            )

            opened = event_row.opened or 0
            replied = event_row.replied or 0
            bounced = event_row.bounced or 0

            results.append(UserPerformanceItem(
                user_id=u.user_id,
                first_name=u.first_name,
                last_name=u.last_name,
                email=u.email,
                role=u.role,
                campaigns_created=campaigns_count,
                total_sent=sent,
                total_opened=opened,
                total_replied=replied,
                total_bounced=bounced,
                open_rate=round((opened / sent) * 100, 2) if sent > 0 else 0.0,
                reply_rate=round((replied / sent) * 100, 2) if sent > 0 else 0.0,
                bounce_rate=round((bounced / sent) * 100, 2) if sent > 0 else 0.0,
            ))

        return results

    def get_dashboard_summary(
        self,
        tenant_id: str,
        start_date: date,
        end_date: date,
        granularity: str = "daily",
        user_id: Optional[str] = None,
    ) -> DashboardSummaryResponse:
        """One call for the full Smartlead-style dashboard."""
        analytics = self.get_global_analytics(tenant_id, start_date, end_date, granularity, user_id)
        provider_perf = self.get_provider_performance(tenant_id, start_date, end_date, user_id)
        top_campaigns = self.get_campaign_comparison(
            tenant_id,
            sort_by="sent_count",
            sort_dir="desc",
            page_size=10,
            start_date=start_date,
            end_date=end_date,
            user_id=user_id,
        )

        # User breakdown only for org-wide view (not when filtering to one user)
        user_breakdown = [] if user_id else self.get_user_breakdown(tenant_id, start_date, end_date)

        return DashboardSummaryResponse(
            kpis=analytics.kpis,
            timeseries=analytics.timeseries,
            lead_stats=self.get_lead_stats(tenant_id, user_id),
            campaign_stats=self.get_campaign_stats(tenant_id, user_id),
            mailbox_stats=self.get_mailbox_stats(tenant_id),
            providers=provider_perf.providers,
            mailbox_health=self.get_mailbox_health(tenant_id, user_id),
            top_campaigns=top_campaigns.items,
            optimization=self.get_optimization_metrics(tenant_id, user_id),
            user_breakdown=user_breakdown,
        )

    # ------------------------------------------------------------------
    # CSV EXPORTS
    # ------------------------------------------------------------------

    def export_global_analytics_csv(
        self, tenant_id: str, start_date: date, end_date: date, user_id: Optional[str] = None
    ) -> io.StringIO:
        """CSV with three sections: overall KPI summary, daily time-series, and per-campaign breakdown."""
        analytics = self.get_global_analytics(tenant_id, start_date, end_date, user_id=user_id)
        campaigns = self.get_campaign_comparison(
            tenant_id,
            sort_by="sent_count",
            sort_dir="desc",
            page_size=10000,
            start_date=start_date,
            end_date=end_date,
            user_id=user_id,
        )

        # Resolve user name for the header label
        user_label = "All Users"
        if user_id:
            u = self.db.query(User).filter(User.user_id == user_id).first()
            if u:
                user_label = f"{u.first_name or ''} {u.last_name or ''}".strip() or u.email

        buf = io.StringIO()
        writer = csv.writer(buf)

        # ── Section 1: Overall KPI Summary ──
        k = analytics.kpis
        writer.writerow(["=== OVERALL SUMMARY ==="])
        writer.writerow(["Period", f"{start_date} to {end_date}"])
        writer.writerow(["Filtered By", user_label])
        writer.writerow(["Active Campaigns", k.active_campaigns])
        writer.writerow(["Total Leads Contacted", k.total_leads_contacted])
        writer.writerow(["Total Emails Sent", k.total_emails_sent])
        writer.writerow(["Total Opened", k.total_opened])
        writer.writerow(["Total Replied", k.total_replied])
        writer.writerow(["Total Bounced", k.total_bounced])
        writer.writerow(["Total Clicked", k.total_clicked])
        writer.writerow(["Overall Open Rate", f"{k.overall_open_rate}%"])
        writer.writerow(["Overall Reply Rate", f"{k.overall_reply_rate}%"])
        writer.writerow(["Overall Bounce Rate", f"{k.overall_bounce_rate}%"])
        writer.writerow([])

        # ── Section 2: Daily Time-Series ──
        writer.writerow(["=== DAILY TIME-SERIES ==="])
        writer.writerow(["Date", "Sent", "Opened", "Replied", "Bounced"])
        for pt in analytics.timeseries:
            writer.writerow([pt.date, pt.sent, pt.opened, pt.replied, pt.bounced])
        writer.writerow([])

        # ── Section 3: Per-Campaign Breakdown ──
        writer.writerow(["=== PER-CAMPAIGN BREAKDOWN ==="])
        writer.writerow([
            "Campaign Name", "Status",
            "Sent", "Opened", "Replied", "Bounced",
            "Open Rate %", "Reply Rate %", "Bounce Rate %",
            "Created At",
        ])
        for c in campaigns.items:
            writer.writerow([
                c.campaign_name,
                c.status,
                c.sent_count,
                c.opened_count,
                c.replied_count,
                c.bounced_count,
                c.open_rate,
                c.reply_rate,
                c.bounce_rate,
                c.created_at.isoformat() if c.created_at else "",
            ])

        buf.seek(0)
        return buf

    def export_campaign_comparison_csv(self, tenant_id: str, user_id: Optional[str] = None) -> io.StringIO:
        """CSV of all campaigns with metrics (optionally scoped to a specific user)."""
        result = self.get_campaign_comparison(tenant_id, page_size=10000, user_id=user_id)

        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "Campaign Name", "Status", "Sent", "Opened", "Replied", "Bounced",
            "Open Rate %", "Reply Rate %", "Bounce Rate %", "Created At",
        ])
        for c in result.items:
            writer.writerow([
                c.campaign_name, c.status, c.sent_count, c.opened_count,
                c.replied_count, c.bounced_count,
                c.open_rate, c.reply_rate, c.bounce_rate,
                c.created_at.isoformat() if c.created_at else "",
            ])
        buf.seek(0)
        return buf

    def export_prospects_csv(self, campaign_id: str) -> io.StringIO:
        """CSV of prospects enrolled in a campaign with engagement data."""

        rows = (
            self.db.query(
                Prospect.email,
                Prospect.first_name,
                Prospect.last_name,
                Prospect.company_name,
                Prospect.designation,
                CampaignProspect.status,
                CampaignProspect.current_step,
            )
            .join(CampaignProspect, CampaignProspect.prospect_id == Prospect.prospect_id)
            .filter(CampaignProspect.campaign_id == campaign_id)
            .all()
        )

        # Get engagement counts per prospect
        engagement = (
            self.db.query(
                EmailMessage.prospect_id,
                func.count(EmailMessage.message_id).label("sent"),
                func.sum(case(
                    (EmailMessage.status == "SENT", 1), else_=0
                )).label("delivered"),
            )
            .filter(EmailMessage.campaign_id == campaign_id)
            .group_by(EmailMessage.prospect_id)
            .all()
        )
        engagement_map = {r.prospect_id: {"sent": r.sent, "delivered": r.delivered} for r in engagement}

        # Get event counts per prospect
        event_counts = (
            self.db.query(
                EmailMessage.prospect_id,
                EmailEvent.event_type,
                func.count(EmailEvent.event_id).label("cnt"),
            )
            .join(EmailEvent, EmailEvent.message_id == EmailMessage.message_id)
            .filter(EmailMessage.campaign_id == campaign_id)
            .group_by(EmailMessage.prospect_id, EmailEvent.event_type)
            .all()
        )
        event_map = {}
        for r in event_counts:
            if r.prospect_id not in event_map:
                event_map[r.prospect_id] = {}
            event_map[r.prospect_id][r.event_type] = r.cnt

        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "Email", "First Name", "Last Name", "Company", "Designation",
            "Status", "Current Step", "Emails Sent", "Opens", "Clicks", "Replies", "Bounces",
        ])
        for r in rows:
            # Match prospect_id from CampaignProspect join
            prospect_row = (
                self.db.query(CampaignProspect.prospect_id)
                .join(Prospect, Prospect.prospect_id == CampaignProspect.prospect_id)
                .filter(
                    CampaignProspect.campaign_id == campaign_id,
                    Prospect.email == r.email,
                )
                .first()
            )
            pid = prospect_row.prospect_id if prospect_row else None
            events = event_map.get(pid, {})

            writer.writerow([
                r.email, r.first_name or "", r.last_name or "",
                r.company_name or "", r.designation or "",
                r.status or "", r.current_step or 0,
                engagement_map.get(pid, {}).get("sent", 0),
                events.get(EmailEvent.EVENT_OPEN, 0),
                events.get(EmailEvent.EVENT_CLICK, 0),
                events.get(EmailEvent.EVENT_REPLY, 0),
                events.get(EmailEvent.EVENT_BOUNCE, 0),
            ])
        buf.seek(0)
        return buf

    # ------------------------------------------------------------------
    # WORD EXPORT HELPERS
    # ------------------------------------------------------------------

    @staticmethod
    def _set_cell_bg(cell, hex_color: str):
        """Fill a table cell background with a hex color (e.g. '2d6bbf')."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    @staticmethod
    def _set_cell_border_bottom(cell, hex_color: str = "E5E7EB"):
        """Add a bottom border to a cell."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:color"), hex_color)
        tcBorders.append(bottom)
        tcPr.append(tcBorders)

    def _add_section_heading(self, doc: Document, text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2D, 0x6B, 0xBF)
        run.font.name = "Calibri"
        # Bottom border on paragraph
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:color"), "2d6bbf")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def _add_kpi_table(self, doc: Document, rows: list[tuple]):
        """Two-column key-value table for KPI summary."""
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, (label, value) in enumerate(rows):
            lc = table.cell(i, 0)
            vc = table.cell(i, 1)
            lc.text = label
            vc.text = str(value)
            lc.paragraphs[0].runs[0].font.size = Pt(10)
            lc.paragraphs[0].runs[0].font.name = "Calibri"
            lc.paragraphs[0].runs[0].bold = True
            lc.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            vc.paragraphs[0].runs[0].font.size = Pt(10)
            vc.paragraphs[0].runs[0].font.name = "Calibri"
            vc.paragraphs[0].runs[0].bold = True
            vc.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            bg = "F9FAFB" if i % 2 == 0 else "FFFFFF"
            self._set_cell_bg(lc, bg)
            self._set_cell_bg(vc, bg)
        # Column widths
        for row in table.rows:
            row.cells[0].width = Inches(2.8)
            row.cells[1].width = Inches(1.5)

    def _add_data_table(self, doc: Document, headers: list[str], data_rows: list[list]):
        """Full data table with blue header row."""
        table = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT

        # Header row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            run = hdr_cells[i].paragraphs[0].runs[0]
            run.bold = True
            run.font.size = Pt(8)
            run.font.name = "Calibri"
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._set_cell_bg(hdr_cells[i], "2D6BBF")

        # Data rows
        for r_idx, row_data in enumerate(data_rows):
            row_cells = table.rows[r_idx + 1].cells
            bg = "F9FAFB" if r_idx % 2 == 0 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].text = str(val) if val is not None else ""
                run = row_cells[c_idx].paragraphs[0].runs[0]
                run.font.size = Pt(8)
                run.font.name = "Calibri"
                run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
                self._set_cell_bg(row_cells[c_idx], bg)

    # ------------------------------------------------------------------
    # WORD EXPORT
    # ------------------------------------------------------------------

    def export_global_analytics_docx(
        self, tenant_id: str, start_date: date, end_date: date, user_id: Optional[str] = None
    ) -> io.BytesIO:
        """
        Full analytics report as a Word document with:
        - Cover header with title and date range
        - Section 1: Overall KPI summary
        - Section 2: Daily time-series table
        - Section 3: Per-campaign breakdown table
        - Section 4: Mailbox health table
        - Section 5: Provider performance table
        """
        analytics = self.get_global_analytics(tenant_id, start_date, end_date, user_id=user_id)
        campaigns = self.get_campaign_comparison(
            tenant_id, sort_by="sent_count", sort_dir="desc",
            page_size=10000, start_date=start_date, end_date=end_date, user_id=user_id,
        )
        mailbox_health = self.get_mailbox_health(tenant_id, user_id)
        provider_perf = self.get_provider_performance(tenant_id, start_date, end_date, user_id=user_id)

        # Resolve user name for report header
        user_label = "All Users"
        if user_id:
            u = self.db.query(User).filter(User.user_id == user_id).first()
            if u:
                user_label = f"{u.first_name or ''} {u.last_name or ''}".strip() or u.email

        doc = Document()

        # ── Page margins ──
        section = doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)

        # ── Cover Header ──
        title_para = doc.add_paragraph()
        title_para.paragraph_format.space_after = Pt(4)
        title_run = title_para.add_run("OUTREACH360 — ANALYTICS REPORT")
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.name = "Calibri"
        title_run.font.color.rgb = RGBColor(0x2D, 0x6B, 0xBF)

        sub_para = doc.add_paragraph()
        sub_para.paragraph_format.space_after = Pt(2)
        sub_run = sub_para.add_run(f"Period: {start_date.strftime('%B %d, %Y')} – {end_date.strftime('%B %d, %Y')}")
        sub_run.font.size = Pt(10)
        sub_run.font.name = "Calibri"
        sub_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        scope_para = doc.add_paragraph()
        scope_para.paragraph_format.space_after = Pt(2)
        scope_run = scope_para.add_run(f"Filtered By: {user_label}")
        scope_run.font.size = Pt(10)
        scope_run.font.name = "Calibri"
        scope_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        gen_para = doc.add_paragraph()
        gen_para.paragraph_format.space_after = Pt(16)
        gen_run = gen_para.add_run(f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
        gen_run.font.size = Pt(9)
        gen_run.font.name = "Calibri"
        gen_run.font.color.rgb = RGBColor(0xA0, 0xAA, 0xB5)

        # ── Section 1: Overall KPI Summary ──
        k = analytics.kpis
        self._add_section_heading(doc, "1. Overall KPI Summary")
        self._add_kpi_table(doc, [
            ("Active Campaigns",     k.active_campaigns),
            ("Total Leads Contacted", k.total_leads_contacted),
            ("Total Emails Sent",    k.total_emails_sent),
            ("Total Opened",         k.total_opened),
            ("Total Replied",        k.total_replied),
            ("Total Bounced",        k.total_bounced),
            ("Total Clicked",        k.total_clicked),
            ("Overall Open Rate",    f"{k.overall_open_rate}%"),
            ("Overall Reply Rate",   f"{k.overall_reply_rate}%"),
            ("Overall Bounce Rate",  f"{k.overall_bounce_rate}%"),
        ])

        # ── Section 2: Daily Time-Series ──
        self._add_section_heading(doc, "2. Daily Sending Activity")
        self._add_data_table(
            doc,
            headers=["Date", "Sent", "Opened", "Replied", "Bounced"],
            data_rows=[
                [pt.date, pt.sent, pt.opened, pt.replied, pt.bounced]
                for pt in analytics.timeseries
            ],
        )

        # ── Section 3: Per-Campaign Breakdown ──
        doc.add_page_break()
        self._add_section_heading(doc, "3. Per-Campaign Breakdown")
        self._add_data_table(
            doc,
            headers=[
                "Campaign Name", "Status", "Sent", "Opened", "Replied",
                "Bounced", "Open Rate %", "Reply Rate %", "Bounce Rate %", "Created",
            ],
            data_rows=[
                [
                    c.campaign_name, c.status,
                    c.sent_count, c.opened_count, c.replied_count, c.bounced_count,
                    f"{c.open_rate}%", f"{c.reply_rate}%", f"{c.bounce_rate}%",
                    c.created_at.strftime("%Y-%m-%d") if c.created_at else "",
                ]
                for c in campaigns.items
            ],
        )

        # ── Section 4: Mailbox Health ──
        self._add_section_heading(doc, "4. Mailbox Health")
        self._add_data_table(
            doc,
            headers=[
                "Mailbox", "Leads Contacted", "Emails Sent",
                "Opened", "Open Rate %", "Replied", "Reply Rate %",
                "Bounced", "Bounce Rate %",
            ],
            data_rows=[
                [
                    m.mailbox, m.lead_contacted, m.email_sent,
                    m.opened, f"{m.opened_rate}%",
                    m.replied, f"{m.replied_rate}%",
                    m.bounced, f"{m.bounce_rate}%",
                ]
                for m in mailbox_health
            ],
        )

        # ── Section 5: Provider Performance ──
        self._add_section_heading(doc, "5. Provider Performance")
        self._add_data_table(
            doc,
            headers=[
                "Provider", "Sent", "Opened", "Replied", "Bounced",
                "Open Rate %", "Reply Rate %", "Bounce Rate %",
            ],
            data_rows=[
                [
                    p.provider, p.sent_count, p.opened_count,
                    p.replied_count, p.bounced_count,
                    f"{p.open_rate}%", f"{p.reply_rate}%", f"{p.bounce_rate}%",
                ]
                for p in provider_perf.providers
            ],
        )

        # ── Footer note ──
        doc.add_paragraph()
        footer_para = doc.add_paragraph()
        footer_run = footer_para.add_run("This report was automatically generated by Outreach360.ai")
        footer_run.font.size = Pt(8)
        footer_run.font.name = "Calibri"
        footer_run.font.color.rgb = RGBColor(0xA0, 0xAA, 0xB5)
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def export_deliverability_csv(self, tenant_id: str) -> io.StringIO:
        """CSV of per-mailbox email health metrics for the tenant."""
        items = self.get_mailbox_health(tenant_id)

        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([
            "Mailbox", "Leads Contacted", "Emails Sent",
            "Opened", "Open Rate %",
            "Replied", "Reply Rate %",
            "Bounced", "Bounce Rate %",
        ])
        for m in items:
            writer.writerow([
                m.mailbox,
                m.lead_contacted,
                m.email_sent,
                m.opened, m.opened_rate,
                m.replied, m.replied_rate,
                m.bounced, m.bounce_rate,
            ])
        buf.seek(0)
        return buf
